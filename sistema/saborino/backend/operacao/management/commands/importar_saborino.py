"""Importa uma planilha mensal da Saborino (.xlsx) para o sistema.

Uso:
    python manage.py importar_saborino "/caminho/Saborino 7_2026.xlsx"
    python manage.py importar_saborino arquivo.xlsx --ano 2026 --mes 7

Idempotente por competência: reimportar o mesmo mês substitui os dados daquele
mês (vendas, pagamentos e compras) sem duplicar. Cadastros (produtos, clientes,
contas, canais) são casados por nome sem diferenciar maiúsculas/minúsculas, o
que corrige a causa raiz dos erros da planilha ("Mousse de Maracuja" vs
"Mousse de maracuja").
"""
import re
from datetime import date
from decimal import Decimal, InvalidOperation

from django.core.management.base import BaseCommand, CommandError
from django.db import transaction

from cadastros.models import Canal, Cliente, Competencia, Conta, Produto
from operacao import services
from operacao.models import Compra, ItemVenda, Venda

FORMA_MAP = {
    'pix': 'PIX', 'dinheiro': 'DINHEIRO', 'cartão': 'CARTAO', 'cartao': 'CARTAO',
}
SOCIO_MAP = {'para nos': 'NOS', 'para nós': 'NOS', 'para o socio 1': 'S1', 'para socio 2': 'S2', 'para o socio 2': 'S2'}
STATUS_PAGO = {'pago'}


def dec(valor, casas='0.01'):
    if valor in (None, ''):
        return None
    try:
        return Decimal(str(valor)).quantize(Decimal(casas))
    except (InvalidOperation, ValueError):
        return None


def limpo(v):
    return str(v).strip() if v not in (None, '') else ''


class Command(BaseCommand):
    help = 'Importa uma planilha mensal (.xlsx) da Saborino.'

    def add_arguments(self, parser):
        parser.add_argument('arquivo', help='Caminho do .xlsx')
        parser.add_argument('--ano', type=int, help='Ano da competência (senão, deduz do nome do arquivo).')
        parser.add_argument('--mes', type=int, help='Mês da competência (senão, deduz do nome do arquivo).')
        parser.add_argument('--dia', type=int, default=15, help='Dia usado nas datas de venda (padrão 15).')

    def handle(self, *args, **opts):
        try:
            import openpyxl
        except ImportError:
            raise CommandError('openpyxl não está instalado. Adicione ao requirements e rebuilde a imagem.')

        caminho = opts['arquivo']
        ano, mes = opts.get('ano'), opts.get('mes')
        if not (ano and mes):
            m = re.search(r'(\d{1,2})[_ -](\d{4})', caminho)
            if not m:
                raise CommandError('Não consegui deduzir ano/mês do nome. Passe --ano e --mes.')
            mes, ano = int(m.group(1)), int(m.group(2))
        if not (1 <= mes <= 12):
            raise CommandError(f'Mês inválido: {mes}')

        try:
            wb = openpyxl.load_workbook(caminho, data_only=True)
        except FileNotFoundError:
            raise CommandError(f'Arquivo não encontrado: {caminho}')

        self.stdout.write(self.style.MIGRATE_HEADING(f'Importando {caminho} -> competência {mes:02d}/{ano}'))
        with transaction.atomic():
            comp = Competencia.para_data(date(ano, mes, 1))
            self._limpar_competencia(comp)
            data_venda = date(ano, mes, opts['dia'])
            rel_v = self._importar_vendas(wb, comp, data_venda)
            rel_c = self._importar_compras(wb, comp)

        self.stdout.write(self.style.SUCCESS(
            f"OK. Vendas: {rel_v['vendas']} (puladas {rel_v['puladas']}), "
            f"pagamentos: {rel_v['pagamentos']}, compras: {rel_c['compras']} (puladas {rel_c['puladas']}). "
            f"Cadastros criados — produtos: {rel_v['produtos']}, clientes: {rel_v['clientes']}, contas: {rel_c['contas']}."
        ))

    def _limpar_competencia(self, comp):
        from operacao.models import Pagamento
        Pagamento.objects.filter(competencia=comp).delete()
        Venda.objects.filter(competencia=comp).delete()  # cascade itens + recebivel
        Compra.objects.filter(competencia=comp).delete()

    # ---- cadastros por nome (case-insensitive) ----
    def _produto(self, nome, cache):
        chave = nome.lower()
        if chave in cache:
            return cache[chave], False
        obj = Produto.objects.filter(nome__iexact=nome).first()
        criado = False
        if not obj:
            obj = Produto.objects.create(nome=nome)
            criado = True
        cache[chave] = obj
        return obj, criado

    def _cliente(self, nome, canal, cache):
        chave = nome.lower()
        if chave in cache:
            return cache[chave], False
        obj = Cliente.objects.filter(nome__iexact=nome).first()
        criado = False
        if not obj:
            obj = Cliente.objects.create(nome=nome, canal=canal)
            criado = True
        cache[chave] = obj
        return obj, criado

    def _canal(self, nome, cache):
        if not nome:
            return None
        from django.utils.text import slugify
        slug = slugify(nome)
        if slug in cache:
            return cache[slug]
        # Casa por slug (ignora acento/caixa): "canal a" -> seed "Canal A".
        obj = Canal.objects.filter(slug=slug).first()
        if not obj:
            obj = Canal.objects.create(nome=nome.title())
        cache[slug] = obj
        return obj

    def _conta(self, nome, cache):
        if not nome:
            return None, False
        chave = nome.lower()
        if chave in cache:
            return cache[chave], False
        obj = Conta.objects.filter(nome__iexact=nome).first()
        criado = False
        if not obj:
            obj = Conta.objects.create(nome=nome)
            criado = True
        cache[chave] = obj
        return obj, criado

    def _importar_vendas(self, wb, comp, data_venda):
        ws = wb['Vendas'] if 'Vendas' in wb.sheetnames else wb.worksheets[1]
        prod_cache, cli_cache, canal_cache, conta_cache = {}, {}, {}, {}
        rel = {'vendas': 0, 'puladas': 0, 'pagamentos': 0, 'produtos': 0, 'clientes': 0}
        for row in ws.iter_rows(min_row=4, max_row=ws.max_row):
            nome = limpo(row[2].value)       # C
            qtd = dec(row[3].value, '0.001')  # D
            valor = dec(row[4].value)         # E
            status = limpo(row[7].value).lower()   # H
            produto_nm = limpo(row[8].value)  # I
            forma = limpo(row[9].value).lower()    # J
            procedencia = limpo(row[11].value)     # L
            destino = limpo(row[12].value).lower() # M
            if not (nome and produto_nm and valor and qtd and qtd > 0):
                if nome or produto_nm:
                    rel['puladas'] += 1
                continue
            canal = self._canal(procedencia, canal_cache)
            cliente, c_new = self._cliente(nome, canal, cli_cache)
            rel['clientes'] += 1 if c_new else 0
            produto, p_new = self._produto(produto_nm, prod_cache)
            rel['produtos'] += 1 if p_new else 0

            preco = (valor / qtd).quantize(Decimal('0.01'))
            venda = Venda.objects.create(
                cliente=cliente, data_venda=data_venda, canal=canal,
                socio_destino=SOCIO_MAP.get(destino, 'NOS'),
            )
            ItemVenda(venda=venda, produto=produto, quantidade=qtd,
                      preco_unitario=preco, subtotal=valor).save()
            venda.recompute_total()
            venda.sincronizar_recebivel()
            rel['vendas'] += 1

            if status in STATUS_PAGO:
                services.quitar_venda(
                    venda=venda, forma_pagamento=FORMA_MAP.get(forma, 'OUTRO'), data=data_venda,
                )
                rel['pagamentos'] += 1
        return rel

    def _importar_compras(self, wb, comp):
        ws = wb['Vendas'] if 'Vendas' in wb.sheetnames else wb.worksheets[1]
        prod_cache, conta_cache = {}, {}
        rel = {'compras': 0, 'puladas': 0, 'contas': 0}
        for row in ws.iter_rows(min_row=4, max_row=ws.max_row):
            desc = limpo(row[14].value)       # O
            qtd = dec(row[15].value, '0.001')  # P
            preco = dec(row[16].value)         # Q
            total = dec(row[17].value)         # R (valor calculado)
            conta_nm = limpo(row[18].value)    # S
            rateio_nm = limpo(row[19].value)   # T
            if not desc:
                continue
            if not total and qtd and preco:
                total = (qtd * preco).quantize(Decimal('0.01'))
            if not total:
                rel['puladas'] += 1
                continue
            conta, ct_new = self._conta(conta_nm, conta_cache)
            rel['contas'] += 1 if ct_new else 0
            produto_rateio = None
            if rateio_nm and rateio_nm.lower() not in ('geral', ''):
                produto_rateio, _ = self._produto(rateio_nm, prod_cache)
            Compra.objects.create(
                competencia=comp, data=date(comp.ano, comp.mes, 15),
                descricao=desc, quantidade=qtd or Decimal('1'),
                preco_unitario=preco or Decimal('0'), valor_total=total,
                conta_origem=conta, produto_rateio=produto_rateio,
            )
            rel['compras'] += 1
        return rel
